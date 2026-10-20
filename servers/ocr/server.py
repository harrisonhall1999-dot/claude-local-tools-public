"""Local document/media extraction MCP server.

Exposes tools so Claude (Desktop / Code) can pull text out of images,
PDFs, Office files, webpages, and audio/video entirely on this machine
instead of spending vision tokens or a full WebFetch call on it.
"""

import os
import tempfile

# Homebrew ffmpeg isn't always on PATH for processes spawned by Claude
# (no login shell), and mlx-whisper shells out to it for audio decoding.
os.environ["PATH"] = "/opt/homebrew/bin:/usr/local/bin:" + os.environ.get("PATH", "")

from mcp.server.mcpserver import MCPServer

mcp = MCPServer("local-tools")

_ocr_engine = None


def _get_engine():
    global _ocr_engine
    if _ocr_engine is None:
        from paddleocr import PaddleOCR

        _ocr_engine = PaddleOCR(
            use_doc_orientation_classify=False,
            use_doc_unwarping=False,
            use_textline_orientation=True,
        )
    return _ocr_engine


def _ocr_image_file(path: str) -> str:
    engine = _get_engine()
    lines = []
    for res in engine.predict(path):
        lines.extend(res.json["res"].get("rec_texts", []))
    return "\n".join(lines)


def _extract_pdf(path: str, dpi: int = 300, min_chars: int = 20) -> str:
    """Per-page hybrid extraction: use the embedded text layer where one
    exists (instant, exact), and only OCR pages that don't have one
    (scanned/image pages, or fully-scanned PDFs)."""
    import pymupdf

    doc = pymupdf.open(path)
    page_texts: list[str | None] = [None] * doc.page_count
    to_ocr: list[tuple[int, str]] = []
    tmp_dir = None
    mat = pymupdf.Matrix(dpi / 72, dpi / 72)

    for i, page in enumerate(doc):
        text = page.get_text().strip()
        if len(text) >= min_chars:
            page_texts[i] = text
        else:
            if tmp_dir is None:
                tmp_dir = tempfile.mkdtemp(prefix="ocr_mcp_pdf_")
            img_path = os.path.join(tmp_dir, f"page_{i + 1:04d}.png")
            page.get_pixmap(matrix=mat).save(img_path)
            to_ocr.append((i, img_path))
    doc.close()

    for i, img_path in to_ocr:
        page_texts[i] = _ocr_image_file(img_path)

    return "\n\n".join(f"--- page {i + 1} ---\n{t}" for i, t in enumerate(page_texts))


@mcp.tool()
def ocr(path: str, dpi: int = 300) -> str:
    """Extract text from a local image or PDF file, entirely on this
    machine (no network calls, no Claude vision tokens spent).

    For PDFs this is a hybrid extractor: pages that already have an
    embedded text layer (digital/exported PDFs, or scanned PDFs that
    were previously OCR'd) are read directly and instantly; pages with
    no text layer (raw scans, images) are rasterized and run through
    local PaddleOCR. Mixed PDFs are handled per-page automatically, so
    this is always the right first call for any PDF or image - no need
    to check for a text layer before calling it.

    Supports common image formats (png, jpg, jpeg, bmp, tiff, webp) and
    multi-page PDFs.

    Args:
        path: Absolute path to the image or PDF file.
        dpi: Rasterization DPI used only for PDF pages that need OCR
            (default 300).

    Returns:
        The extracted text. For multi-page PDFs, pages are separated by
        a "--- page N ---" marker.
    """
    path = os.path.expanduser(path)
    if not os.path.isfile(path):
        return f"Error: file not found: {path}"

    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return _extract_pdf(path, dpi=dpi)

    return _ocr_image_file(path)


def _read_docx(path: str) -> str:
    import docx

    doc = docx.Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _read_pptx(path: str) -> str:
    from pptx import Presentation

    prs = Presentation(path)
    chunks = []
    for i, slide in enumerate(prs.slides, 1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs)
                    if text.strip():
                        lines.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    lines.append("\t".join(cell.text for cell in row.cells))
        chunks.append(f"--- slide {i} ---\n" + "\n".join(lines))
    return "\n\n".join(chunks)


def _read_xlsx(path: str) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    chunks = []
    for ws in wb.worksheets:
        lines = [f"--- sheet: {ws.title} ---"]
        for row in ws.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                lines.append("\t".join("" if c is None else str(c) for c in row))
        chunks.append("\n".join(lines))
    return "\n\n".join(chunks)


@mcp.tool()
def read_document(path: str) -> str:
    """Extract text from a local Word/PowerPoint/Excel file.

    Reads .docx/.pptx/.xlsx directly from the file's XML - no rendering,
    no vision tokens. Word: paragraphs + table cells. PowerPoint: text
    per slide (with a "--- slide N ---" marker) including tables. Excel:
    all sheets as tab-separated rows (with a "--- sheet: NAME ---"
    marker), formulas resolved to their last-calculated value.

    Legacy binary formats (.doc, .ppt, .xls) are not supported.

    Args:
        path: Absolute path to the .docx, .pptx, or .xlsx file.

    Returns:
        The extracted text.
    """
    path = os.path.expanduser(path)
    if not os.path.isfile(path):
        return f"Error: file not found: {path}"

    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        return _read_docx(path)
    if ext == ".pptx":
        return _read_pptx(path)
    if ext == ".xlsx":
        return _read_xlsx(path)
    return f"Error: unsupported extension '{ext}' (supported: .docx, .pptx, .xlsx)"


@mcp.tool()
def fetch_url(url: str) -> str:
    """Fetch a webpage and return its main readable text.

    Runs entirely on this machine: a plain HTTP GET followed by local
    boilerplate stripping (nav/ads/footers removed, main article/body
    text and tables kept). Does NOT execute JavaScript, so pages that
    render their content client-side may return little or nothing - for
    those, a normal WebFetch call is still needed.

    The extractor trims by guesswork and can drop parts of a page (it has
    dropped statute text and block lists). For a legal source, or anything
    to be saved, quoted or cited, use the files server's to_markdown instead.

    Args:
        url: The page URL to fetch (http/https).

    Returns:
        The extracted readable text, or an error message if nothing
        extractable was found.
    """
    import requests
    import trafilatura

    try:
        resp = requests.get(url, timeout=20, headers={"User-Agent": "Mozilla/5.0"})
        resp.raise_for_status()
    except requests.RequestException as e:
        return f"Error fetching {url}: {e}"

    text = trafilatura.extract(
        resp.text, url=url, include_comments=False, include_tables=True
    )
    if not text:
        return (
            f"Error: no readable content extracted from {url} "
            "(page may require JavaScript to render)."
        )
    return text


TRANSCRIBE_MODEL = "mlx-community/whisper-large-v3-turbo"

# Speaker diarization (who spoke when) with sherpa-onnx: a pyannote
# segmentation model plus an English speaker-embedding model, ~50 MB in all,
# downloaded on first use. name -> (release URL, file to use once unpacked).
_SHERPA_RELEASES = "https://github.com/k2-fsa/sherpa-onnx/releases/download"
SPEAKER_MODELS = {
    "segmentation": (
        f"{_SHERPA_RELEASES}/speaker-segmentation-models/sherpa-onnx-pyannote-segmentation-3-0.tar.bz2",
        "sherpa-onnx-pyannote-segmentation-3-0/model.onnx"),
    "embedding": (
        f"{_SHERPA_RELEASES}/speaker-recongition-models/nemo_en_titanet_small.onnx",
        "nemo_en_titanet_small.onnx"),
}
SPEAKER_MODEL_DIR = os.path.expanduser("~/.cache/sherpa-onnx")
# Clustering distance for telling voices apart; sherpa's default (0.5) split
# one host on a Zoom recording into a dozen speakers, 0.9 found the real five.
SPEAKER_THRESHOLD = 0.9
# "Speakers" with less total speech than this are noise fragments; their turns
# go to the nearest real speaker.
MIN_SPEAKER_SECONDS = 5.0
# A transcript block ends at a speaker change, a pause this long, or at the
# end of a sentence once it has run this long.
PAUSE_BREAK = 2.0
BLOCK_SECONDS = 90.0


def _drop_repeats(segments, keep=2):
    """Drop Whisper segments that repeat the one before more than `keep` times
    in a row, the telltale of a hallucination loop over silence or noise."""
    kept, prev, run = [], None, 0
    for seg in segments:
        key = seg["text"].strip().lower()
        run = run + 1 if key == prev else 1
        prev = key
        if run <= keep:
            kept.append(seg)
    return kept


def _speaker_model(name):
    """Path to a diarization model, downloading it the first time."""
    url, rel = SPEAKER_MODELS[name]
    path = os.path.join(SPEAKER_MODEL_DIR, rel)
    if not os.path.exists(path):
        import tarfile
        import urllib.request

        os.makedirs(SPEAKER_MODEL_DIR, exist_ok=True)
        tmp, _ = urllib.request.urlretrieve(url)
        try:
            if url.endswith(".tar.bz2"):
                with tarfile.open(tmp) as tar:
                    tar.extractall(SPEAKER_MODEL_DIR, filter="data")
            else:
                os.replace(tmp, path)
        finally:
            if os.path.exists(tmp):
                os.remove(tmp)
    return path


def _diarize(audio, num_speakers=0):
    """Return [(start, end, speaker)] turns for 16 kHz mono float32 audio.
    num_speakers=0 lets the clustering decide how many there are."""
    import sherpa_onnx

    config = sherpa_onnx.OfflineSpeakerDiarizationConfig(
        segmentation=sherpa_onnx.OfflineSpeakerSegmentationModelConfig(
            pyannote=sherpa_onnx.OfflineSpeakerSegmentationPyannoteModelConfig(
                model=_speaker_model("segmentation")),
            num_threads=4),
        embedding=sherpa_onnx.SpeakerEmbeddingExtractorConfig(
            model=_speaker_model("embedding"), num_threads=4),
        clustering=sherpa_onnx.FastClusteringConfig(
            num_clusters=num_speakers if num_speakers > 0 else -1,
            threshold=SPEAKER_THRESHOLD),
    )
    result = sherpa_onnx.OfflineSpeakerDiarization(config).process(audio)
    turns = [(t.start, t.end, t.speaker) for t in result.sort_by_start_time()]

    total = {}
    for start, end, speaker in turns:
        total[speaker] = total.get(speaker, 0.0) + end - start
    real = [t for t in turns if total[t[2]] >= MIN_SPEAKER_SECONDS]
    if not real or len(real) == len(turns):
        return turns
    return [t if total[t[2]] >= MIN_SPEAKER_SECONDS
            else (t[0], t[1], min(real, key=lambda r: abs(r[0] - t[0]))[2])
            for t in turns]


def _speaker_at(start, end, turns, max_gap=3.0):
    """The speaker who talks most during [start, end], else the nearest one
    within max_gap seconds. None means no one is talking near there, so the
    line is likely Whisper hallucinating over silence."""
    best, best_overlap = None, 0.0
    for t_start, t_end, speaker in turns:
        overlap = min(end, t_end) - max(start, t_start)
        if overlap > best_overlap:
            best, best_overlap = speaker, overlap
    if best is None and turns:
        gap, best = min((max(t_start - end, start - t_end), speaker)
                        for t_start, t_end, speaker in turns)
        if gap > max_gap:
            best = None
    return best


def _timestamp(seconds):
    s = int(seconds)
    return f"{s // 3600:02d}:{s % 3600 // 60:02d}:{s % 60:02d}"


def _format_transcript(segments, turns=None):
    """Lay out Whisper segments as timestamped blocks, one per speaker turn
    (or per paragraph without speakers), with speakers numbered by first
    appearance: "[00:01:23] Speaker 1: ..."."""
    # Speakers are matched per segment (a sentence or phrase), not per word, so
    # a turn boundary that diarization places mid-sentence doesn't split it.
    pieces = []  # (start, end, text)
    for seg in segments:
        words = seg.get("words") or []
        start = words[0]["start"] if words else seg["start"]
        end = words[-1]["end"] if words else seg["end"]
        pieces.append((start, end, seg["text"]))

    blocks, names = [], {}
    for start, end, text in pieces:
        speaker = None
        if turns:
            raw = _speaker_at(start, end, turns)
            if raw is None:
                continue
            speaker = names.setdefault(raw, f"Speaker {len(names) + 1}")
        if blocks:
            b = blocks[-1]
            long_block = end - b["start"] > BLOCK_SECONDS and b["text"].rstrip().endswith((".", "?", "!"))
            if speaker == b["speaker"] and start - b["end"] < PAUSE_BREAK and not long_block:
                b["text"] += text
                b["end"] = end
                continue
        blocks.append({"start": start, "end": end, "speaker": speaker, "text": text})

    lines = []
    for b in blocks:
        label = f" {b['speaker']}:" if b["speaker"] else ""
        lines.append(f"[{_timestamp(b['start'])}]{label} {b['text'].strip()}")
    return "\n\n".join(lines)


def _split_list(value):
    return [v.strip() for v in value.split(",") if v.strip()]


def _fix_spellings(text, terms):
    """Respell capitalized words that are near-misses of the user's names and
    terms ("Tonya" -> "Tanya"). Whisper only sees a prompt for its first 30
    seconds, so this is what keeps names right for the rest of a recording."""
    import difflib
    import re

    targets = {}
    for term in terms:
        for word in re.findall(r"[A-Za-z][\w'-]{3,}", term):
            if word[0].isupper():
                targets[word.lower()] = word

    def fix(m):
        word = m.group(0)
        if word.lower() in targets:
            return targets[word.lower()] if word.isupper() == targets[word.lower()].isupper() else word
        close = difflib.get_close_matches(word.lower(), targets, n=1, cutoff=0.8)
        if close and close[0][0] == word[0].lower():
            return targets[close[0]]
        return word

    return re.sub(r"\b[A-Z][\w'-]{3,}", fix, text)


def _speaker_summary(text, snippets=2, chars=140):
    """Per speaker: words spoken, first time heard, and the start of their
    first few turns, so names can be matched without reading the whole file."""
    import re

    info = {}
    for m in re.finditer(r"^\[([\d:]+)\] ([^:\n]+): (.*)$", text, re.M):
        when, who, said = m.groups()
        i = info.setdefault(who, {"first": when, "words": 0, "lines": []})
        i["words"] += len(said.split())
        if len(i["lines"]) < snippets:
            i["lines"].append(said[:chars] + ("..." if len(said) > chars else ""))
    out = []
    for who, i in info.items():
        out.append(f"{who} ({i['words']} words, first at {i['first']}):")
        out.extend(f"  - {line}" for line in i["lines"])
    return "\n".join(out)


@mcp.tool()
def rename_speakers(path: str, names: dict[str, str]) -> str:
    """Replace "Speaker N" labels in a transcript file made by transcribe.

    Use after transcribe, once you've matched speakers to names from its
    per-speaker summary (who is greeted or handed the mic, self-introductions,
    roles the user described). Only rename speakers you're confident about;
    leave the rest numbered.

    Args:
        path: Absolute path to the transcript .txt file.
        names: Map of current label to name, e.g. {"Speaker 4": "Tatiana
            (host)", "Speaker 8": "Tanya Richardson"}. Two labels may map to
            the same name if one person was split in two.

    Returns:
        How many lines were relabelled, and the updated speaker summary.
    """
    import re

    path = os.path.expanduser(path)
    if not os.path.isfile(path):
        return f"Error: file not found: {path}"
    with open(path, encoding="utf-8") as f:
        text = f.read()
    count = 0

    def relabel(m):
        nonlocal count
        new = names.get(m.group(2))
        if new is None:
            return m.group(0)
        count += 1
        return f"{m.group(1)}{new}:"

    text = re.sub(r"^(\[[\d:]+\] )([^:\n]+):", relabel, text, flags=re.M)
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)
    return f"Relabelled {count} turns in {path}.\n\n{_speaker_summary(text, snippets=1, chars=80)}"


@mcp.tool()
def transcribe(path: str, model: str = TRANSCRIBE_MODEL, output_path: str = "",
               speakers: bool = True, num_speakers: int = 0,
               speaker_names: str = "", context: str = "", language: str = "") -> str:
    """Transcribe a local audio or video file to a readable transcript.

    BEFORE CALLING, ask the user in one short message for whatever they
    haven't already told you (any of it may be "don't know"; then use the
    defaults):
      - how many people speak (-> num_speakers),
      - their names, and roles if useful, e.g. "Tatiana (host), Tanya
        Richardson (guest)" (-> speaker_names),
      - the topic and any names, acronyms or jargon likely to come up
        (-> context),
      - the language, if not English (-> language),
      - whether they need speaker labels at all (-> speakers),
      - where to save it (-> output_path; default: next to the audio, same
        name, .txt).
    This is worth one round trip: it improves speaker separation and name
    spelling, and lets you name the speakers without reading the transcript.

    Runs entirely on this machine using Whisper via MLX (Metal-
    accelerated on Apple Silicon) - no network calls, no Claude usage
    spent listening to the file. Video files have their audio track
    decoded automatically (via ffmpeg). Accepts common formats: mp3,
    wav, m4a, mp4, mov, etc.

    The transcript comes in timestamped blocks, one per speaker turn:
    "[00:01:23] Speaker 1: ...". Speakers are told apart by voice
    (sherpa-onnx, offline) and numbered in order of first appearance; it
    can't know names, so infer those from the conversation if needed. Set
    num_speakers when you know how many people talk, which helps accuracy.
    Set speakers=False for a single speaker or a faster run (blocks then
    break at pauses instead).

    Uses ~2-3 GB of RAM while running (released afterwards). Give the user
    a one-line heads-up with a reminder to close Teams, Outlook and extra
    Chrome tabs first. An hour of audio takes about 11 minutes: ~8 for
    Whisper and ~3-4 for telling speakers apart (speakers=False skips that).

    For anything longer than a few minutes, set output_path so the
    transcript is saved to a file instead of returned. An hour of speech is
    ~10k words, which you usually don't need in context. Then read only the
    parts you need, or summarize the file with the local_llm tool.

    Repeated lines (Whisper's habit of looping "Thank you." over silence)
    are cut to two, and with speakers on, lines far from any detected speech
    are dropped. A stray line can still slip through, usually at the end.

    Args:
        path: Absolute path to the audio or video file.
        model: Hugging Face MLX Whisper model repo. Default is
            large-v3-turbo (accurate on lectures, accents and legal
            vocabulary). Use "mlx-community/whisper-base-mlx" for a much
            faster, rougher pass.
        output_path: Optional path for a .txt file to save the transcript to.
        speakers: Label who is speaking (default True).
        num_speakers: How many people speak, if known; 0 guesses.
        speaker_names: Comma-separated names (roles in brackets are fine).
            Used to spell names right; labels stay "Speaker N" until you
            call rename_speakers.
        context: Topic plus names, acronyms and terms, comma-separated.
            Primes Whisper and fixes near-miss spellings of capitalized
            terms throughout.
        language: Language code such as "en" or "es"; empty detects it.

    Returns:
        The transcript, or, with output_path, where it was saved and how
        long it is. With speakers, also a short per-speaker summary (words,
        first time heard, opening lines): match it to speaker_names and call
        rename_speakers, instead of reading the transcript to find out who
        is who.
    """
    import re

    import mlx.core as mx
    import mlx_whisper
    import numpy as np
    from mlx_whisper.audio import load_audio
    from mlx_whisper.transcribe import ModelHolder

    path = os.path.expanduser(path)
    if not os.path.isfile(path):
        return f"Error: file not found: {path}"

    # Use the downloaded copy if there is one, so this works offline and
    # doesn't check Hugging Face on every call; download only on first use.
    from huggingface_hub import snapshot_download

    try:
        model_dir = snapshot_download(model, local_files_only=True)
    except Exception:
        model_dir = snapshot_download(model)

    try:
        # 16 kHz mono float32, shared with diarization. load_audio returns an
        # MLX array, which sherpa-onnx would read one sample at a time
        # (~10x slower than the model itself), so hand both a NumPy copy.
        audio = np.array(load_audio(path), dtype=np.float32)
    except Exception as e:
        return f"Error: couldn't decode audio from {path}: {e}"

    try:
        # Not conditioning on the previous text stops one hallucinated line
        # (e.g. "Thank you." over trailing silence) from repeating to the end.
        # Word timestamps give tighter segment times for matching speakers,
        # and let Whisper skip hallucinations over long silences.
        names = _split_list(speaker_names)
        prompt = " ".join(filter(None, [
            context.strip(), f"Speakers: {', '.join(names)}." if names else ""]))
        result = mlx_whisper.transcribe(
            audio, path_or_hf_repo=model_dir, condition_on_previous_text=False,
            word_timestamps=True, hallucination_silence_threshold=2.0,
            initial_prompt=prompt or None, language=language.strip() or None)
    finally:
        # mlx_whisper keeps the model loaded for the life of this process;
        # give the RAM back, since this Mac only has 16 GB.
        ModelHolder.model = None
        ModelHolder.model_path = None
        mx.clear_cache()
    segments = _drop_repeats(result["segments"])

    # Diarize after Whisper's memory is released, to keep the peak low.
    turns, note = None, ""
    if speakers:
        try:
            turns = _diarize(audio, num_speakers)
        except Exception as e:
            note = f" Speakers not labelled ({e!r})."
    text = _format_transcript(segments, turns)
    terms = [re.sub(r"\(.*?\)", "", n) for n in names] + _split_list(context)
    if terms:
        text = _fix_spellings(text, terms)
    summary = f"\n\n{_speaker_summary(text)}" if turns else ""

    if output_path:
        output_path = os.path.expanduser(output_path)
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(text + "\n")
        found = f", {len(set(t[2] for t in turns))} speakers" if turns else ""
        return (f"Saved transcript ({len(text.split())} words{found}) to "
                f"{output_path}.{note}{summary}")
    return text + (f"\n\n({note.strip()})" if note else "")


if __name__ == "__main__":
    mcp.run()
